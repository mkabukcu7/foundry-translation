# SPDX-License-Identifier: MIT
import io

from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from services import docintel_client, slide_render


def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from all_shapes(shape.shapes)


def extract(data: bytes) -> dict:
    presentation = Presentation(io.BytesIO(data))
    has_images = any(shape.shape_type == 13 for slide in presentation.slides for shape in all_shapes(slide.shapes))
    if has_images:
        # Use rendered OCR for the complete slide rather than double-counting native text.
        rendered = PdfReader(io.BytesIO(slide_render.render(data)))
        if len(rendered.pages) != len(presentation.slides):
            raise ValueError("Rendered slide count differs from the presentation")
        segments = []
        for i, page in enumerate(rendered.pages):
            writer, output = PdfWriter(), io.BytesIO()
            writer.add_page(page)
            writer.write(output)
            result = docintel_client.extract(output.getvalue(), "prebuilt-read", allow_empty=True)
            segments.append({
                "text": "\n".join(p["text"] for p in result["segments"]),
                "role": "slide", "location": f"slide:{i + 1}",
            })
        return {"segments": segments, "language": ""}
    segments = []
    for i, slide in enumerate(presentation.slides):
        text = []
        for shape in all_shapes(slide.shapes):
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
            if shape.has_table:
                text.extend(cell.text for row in shape.table.rows for cell in row.cells)
        segments.append({"text": "\n".join(text), "role": "slide", "location": f"slide:{i + 1}"})
    return {"segments": segments, "language": ""}
