# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.
"""Generate synthetic Italian/Spanish test documents and samples/expected.json.

All names, numbers and addresses are fictitious.
Run: python samples/generate_samples.py
"""
import io
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).parent / "out"

# Seeded values the translation must preserve (checked later by Pass A).
CLAIM_IT = {
    "policy_numbers": ["POL-2024-558812"],
    "claim_numbers": ["SIN-2025-004471"],
    "amounts": ["€ 12.450,75", "€ 1.200,00", "€ 11.250,75"],
    "dates": ["14/03/2025", "02/04/2025"],
    "names": ["Giulia Rossetti", "Alberto Ferraro"],
    "addresses": ["Via Garibaldi 42, 20121 Milano"],
}
SCAN_ES = {
    "policy_numbers": ["POL-2024-771203"],
    "claim_numbers": ["SIE-2025-009318"],
    "amounts": ["8.320,50 €"],
    "dates": ["03/02/2025"],
    "names": ["Carmen Ortega Vidal"],
    "addresses": ["Calle Mayor 15, 28013 Madrid"],
}
PPTX_ES = {
    "policy_numbers": ["POL-2024-330145"],
    "claim_numbers": ["SIE-2025-001207"],
    "amounts": ["5.600,00 €"],
    "dates": ["21/01/2025"],
    "names": ["Javier Morales Peña"],
    "addresses": [],
}


FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def font(size):
    # The built-in Pillow font lacks accents and the euro sign, so a TTF is required.
    for p in FONT_CANDIDATES:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    raise SystemExit("No TTF font found; add a path to FONT_CANDIDATES.")


def claim_form_pdf(path):
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    rows = [
        ["Voce", "Descrizione", "Importo"],
        ["1", "Danni da acqua - cucina", "€ 7.450,75"],
        ["2", "Danni da acqua - soggiorno", "€ 5.000,00"],
        ["", "Totale richiesto", "€ 12.450,75"],
        ["", "Franchigia", "€ 1.200,00"],
        ["", "Indennizzo liquidabile", "€ 11.250,75"],
    ]
    table = Table(rows, colWidths=[40, 300, 100])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
    ]))
    doc.build([
        Paragraph("Modulo di denuncia di sinistro", styles["Title"]),
        Paragraph("Numero di polizza: POL-2024-558812", styles["Normal"]),
        Paragraph("Numero di sinistro: SIN-2025-004471", styles["Normal"]),
        Paragraph("Contraente: Giulia Rossetti, Via Garibaldi 42, 20121 Milano", styles["Normal"]),
        Paragraph("Data del sinistro: 14/03/2025. Data di denuncia: 02/04/2025.", styles["Normal"]),
        Spacer(1, 12),
        Paragraph(
            "Il contraente dichiara che in data 14/03/2025 una perdita dalla tubazione "
            "dell'appartamento al piano superiore ha causato danni da acqua alla cucina e "
            "al soggiorno. Il perito Alberto Ferraro ha effettuato il sopralluogo. "
            "La copertura prevede una franchigia fissa e il massimale della polizza "
            "non risulta superato.",
            styles["Normal"],
        ),
        Spacer(1, 12),
        table,
    ])


def scanned_pdf(path):
    """Image-only PDF: all text is rasterised, so there is no text layer."""
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    lines = [
        ("Notificación de siniestro", 56),
        ("Número de póliza: POL-2024-771203", 36),
        ("Número de siniestro: SIE-2025-009318", 36),
        ("Asegurado: Carmen Ortega Vidal", 36),
        ("Domicilio: Calle Mayor 15, 28013 Madrid", 36),
        ("Fecha del siniestro: 03/02/2025", 36),
        ("Importe reclamado: 8.320,50 €", 36),
        ("El asegurado comunica un robo con fuerza en su vivienda habitual.", 32),
        ("La franquicia aplicable se descontará de la indemnización.", 32),
    ]
    y = 120
    for text, size in lines:
        d.text((100, y), text, fill="black", font=font(size))
        y += size + 40
    img = img.rotate(0.6, fillcolor="white")  # slight skew, scan-like
    img.save(path, "PDF", resolution=150.0)


def pptx_with_screenshot(path):
    shot = Image.new("RGB", (900, 360), (245, 245, 250))
    d = ImageDraw.Draw(shot)
    d.rectangle([0, 0, 900, 50], fill=(30, 70, 140))
    d.text((20, 10), "Sistema de siniestros - Detalle", fill="white", font=font(28))
    for i, t in enumerate([
        "Póliza: POL-2024-330145",
        "Siniestro: SIE-2025-001207",
        "Fecha de ocurrencia: 21/01/2025",
        "Importe estimado: 5.600,00 €",
    ]):
        d.text((30, 90 + i * 60), t, fill="black", font=font(30))
    buf = io.BytesIO()
    shot.save(buf, "PNG")
    buf.seek(0)

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Resumen del siniestro"
    tf = s1.placeholders[1].text_frame
    tf.text = "Asegurado: Javier Morales Peña"
    for line in ["Póliza: POL-2024-330145", "Fecha del siniestro: 21/01/2025",
                 "Importe estimado: 5.600,00 €"]:
        tf.add_paragraph().text = line
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Captura del sistema"
    s2.shapes.add_picture(buf, Inches(0.6), Inches(1.8), width=Inches(8.8))
    box = s2.shapes.add_textbox(Inches(0.6), Inches(6.3), Inches(8.8), Inches(0.6))
    box.text_frame.text = "La cobertura incluye daños por robo según condiciones particulares."
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    prs.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    claim_form_pdf(OUT / "claim_form_it.pdf")
    scanned_pdf(OUT / "scanned_claim_es.pdf")
    pptx_with_screenshot(OUT / "claim_summary_es.pptx")
    expected = {
        "claim_form_it.pdf": {"sourceLanguage": "it", **CLAIM_IT},
        "scanned_claim_es.pdf": {"sourceLanguage": "es", **SCAN_ES},
        "claim_summary_es.pptx": {"sourceLanguage": "es", **PPTX_ES},
    }
    (Path(__file__).parent / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print("Wrote", *sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
