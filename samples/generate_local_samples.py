# SPDX-License-Identifier: MIT
"""Generate small Italian samples for the storage-free workflow."""

from PIL import Image, ImageDraw
from pptx import Presentation
from reportlab.pdfgen import canvas

from samples.generate_samples import OUT, font

LINES = [
    "Denuncia di sinistro",
    "Numero di polizza: POL-2024-558812",
    "Numero di sinistro: SIN-2025-004471",
    "Contraente: Giulia Rossetti",
    "La polizza copre i danni da acqua.",
    "La franchigia e di EUR 500,00.",
]


def main():
    OUT.mkdir(exist_ok=True)
    pdf = canvas.Canvas(str(OUT / "local_claim_it.pdf"))
    for i, line in enumerate(LINES):
        pdf.drawString(60, 780 - i * 35, line)
    pdf.save()
    scan = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(scan)
    for i, line in enumerate(LINES):
        draw.text((80, 120 + i * 80), line, fill="black", font=font(36))
    scan.save(OUT / "local_scanned_it.pdf", "PDF", resolution=150)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = LINES[0]
    slide.placeholders[1].text = "\n".join(LINES[1:])
    presentation.save(OUT / "local_claim_it.pptx")
    print("Generated 3 synthetic Italian documents")


if __name__ == "__main__":
    main()
