# SPDX-License-Identifier: MIT
"""Storage-free Translator adapters. No retries for billable translation requests."""

import io
from urllib.parse import urlparse

import requests
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from services.config import ExternalServiceError, credential, required
from services.pptx_extract import all_shapes

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def native_pptx_only(data: bytes):
    presentation = Presentation(io.BytesIO(data))
    for slide in presentation.slides:
        layout = slide.slide_layout
        parts = (slide.part, layout.part, layout.slide_master.part)
        if any(part._element.xpath(".//a:blip") for part in parts):
            raise ValueError("Storage-free PPTX mode supports native text only; use a scanned PDF for the OCR demo")
        shapes = (slide.shapes, layout.shapes, layout.slide_master.shapes)
        if any(shape.shape_type == 13 or (
            shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        ) for group in shapes for shape in all_shapes(group)):
            raise ValueError("Storage-free PPTX mode supports native text only; use a scanned PDF for the OCR demo")


def request(path, job_id, **kwargs):
    endpoint = required("TRANSLATOR_ENDPOINT").rstrip("/")
    parsed = urlparse(endpoint)
    if parsed.scheme != "https" or not parsed.netloc or parsed.path not in ("", "/"):
        raise ValueError("TRANSLATOR_ENDPOINT must be an HTTPS resource origin")
    response = requests.post(
        endpoint + path,
        headers={
            "Authorization": f"Bearer {credential().get_token('https://cognitiveservices.azure.com/.default').token}",
            "x-ms-client-request-id": job_id,
        },
        timeout=(10, 180), allow_redirects=False, **kwargs,
    )
    if response.status_code != 200:
        raise ExternalServiceError(f"Translator HTTP {response.status_code}; jobId={job_id}")
    return response


def translate_pptx(data, glossary_data, source_language, target_language, job_id):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Synchronous document translation supports documents up to 10 MB")
    if len(glossary_data) > 1024 * 1024:
        raise ValueError("Synchronous glossary exceeds 1 MB")
    native_pptx_only(data)
    files = {"document": ("document.pptx", data, PPTX_TYPE)}
    if glossary_data:
        files["glossary"] = ("glossary.tsv", glossary_data, "text/tab-separated-values")
    response = request("/translator/document:translate", job_id, params={
        "api-version": "2024-05-01", "sourceLanguage": source_language, "targetLanguage": target_language,
    }, files=files)
    if not response.content.startswith(b"PK"):
        raise ExternalServiceError(f"Translator did not return a PPTX; jobId={job_id}")
    native_pptx_only(response.content)
    return response.content, {"mode": "synchronous-document", "glossaryApplied": bool(glossary_data),
                              "charactersCharged": None, "imagesTranslated": 0}


def translate_segments(segments, source_language, target_language, job_id):
    translated, submitted, charged = [], 0, []
    for segment in segments:
        text, output = segment["text"], []
        # Bound each request below the 50,000-character limit, without overlap or lost text.
        for start in range(0, len(text), 10000):
            part = text[start:start + 10000]
            response = request("/translator/text/v3.0/translate", job_id, params={
                "api-version": "3.0", "from": source_language, "to": target_language,
            }, json=[{"Text": part}])
            try:
                body = response.json()
            except ValueError:
                raise ExternalServiceError(f"Translator returned invalid JSON; jobId={job_id}") from None
            if (not isinstance(body, list) or len(body) != 1
                    or not isinstance(body[0], dict)
                    or not isinstance(body[0].get("translations"), list)
                    or len(body[0]["translations"]) != 1
                    or not isinstance(body[0]["translations"][0], dict)
                    or body[0]["translations"][0].get("to") != target_language
                    or not isinstance(body[0]["translations"][0].get("text"), str)):
                raise ExternalServiceError(f"Translator returned unexpected text result; jobId={job_id}")
            output.append(body[0]["translations"][0]["text"])
            submitted += len(part)
            metered = response.headers.get("X-Metered-Usage")
            if metered is not None and not metered.isdecimal():
                raise ExternalServiceError(f"Translator returned invalid metered usage; jobId={job_id}")
            charged.append(int(metered) if metered is not None else None)
        translated.append({**segment, "text": "".join(output)})
    return {"segments": translated, "language": target_language}, {
        "mode": "ocr-then-text", "glossaryApplied": False,
        "charactersSubmitted": submitted,
        "charactersCharged": sum(charged) if all(c is not None for c in charged) else None,
        "imagesTranslated": 0,
    }
