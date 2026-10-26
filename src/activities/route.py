# SPDX-License-Identifier: MIT
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from services import jobs_table, storage

WATERMARK = "Machine translation \u2013 requires review"


def watermark(data: bytes, filename: str) -> bytes:
    if filename.lower().endswith(".pdf"):
        document, writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data))), PdfWriter()
        for page in document.pages:
            page.transfer_rotation_to_content()
            width, height = float(page.mediabox.width), float(page.mediabox.height) + 28
            new = writer.add_blank_page(width=width, height=height)
            new.merge_page(page)
            header = io.BytesIO()
            drawing = canvas.Canvas(header, pagesize=(width, height))
            drawing.setFont("Helvetica", 10)
            drawing.setFillColorRGB(0.7, 0, 0)
            drawing.drawString(12, height - 16, WATERMARK)
            drawing.save()
            new.merge_page(PdfReader(header).pages[0])
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()
    presentation = Presentation(io.BytesIO(data))
    old_height = presentation.slide_height
    presentation.slide_height = old_height + Inches(0.4)
    for slide in presentation.slides:
        shape = slide.shapes.add_textbox(Inches(0.15), old_height, presentation.slide_width, Inches(0.4))
        shape.text_frame.text = WATERMARK
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.color.rgb = Pt(12), RGBColor(180, 0, 0)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def run(job):
    report = storage.get_json(job["jobId"], "report")
    name, data = storage.target(job)
    container = f"target-{job['jobId']}"
    if report["decision"] == "flagged":
        # Keep Translator output immutable so a retried route activity never adds a second watermark.
        data = watermark(data, name)
        container, name = "reports", f"{job['jobId']}/reviewed.{job['fileName'].rsplit('.', 1)[-1]}"
        storage.put(container, name, data)
    severity_order = {"critical": 0, "major": 1, "minor": 2}
    reasons = [i["explanation"] for i in sorted(report["issues"], key=lambda i: severity_order[i["severity"]])]
    reasons.extend(f"Pass A: unmatched {f['type']} at {f['location']}" for f in report["passA"]
                   if f["match"] == "unmatched")
    report["topReasons"] = reasons[:5]
    current = jobs_table.get(job["jobId"])
    report["costSignals"]["durationsSeconds"] = {
        key.removeprefix("duration_"): value for key, value in current.items() if key.startswith("duration_")
    }
    storage.put_json(job["jobId"], "report", report)
    # Keep reasons/excerpts in blobs, not Table Storage's small entity properties.
    jobs_table.update(job["jobId"], score=report["score"], decision=report["decision"],
                      reportPath=f"{job['jobId']}/report.json", resultPath=name, resultContainer=container,
                      tokens=report["costSignals"]["tokens"])
    return {"status": "complete" if report["decision"] == "passed" else "needs_review"}
