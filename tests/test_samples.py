import io

from pptx import Presentation
from pypdf import PdfReader

from samples.generate_samples import claim_form_pdf, pptx_with_screenshot, scanned_pdf
from services import docintel_client, pptx_extract, slide_render
from tests.test_pipeline import pdf


def test_synthetic_samples_cover_text_tables_and_image_only_content(tmp_path):
    native, scan, slides = [tmp_path / name for name in ("native.pdf", "scan.pdf", "slides.pptx")]
    claim_form_pdf(native)
    scanned_pdf(scan)
    pptx_with_screenshot(slides)
    assert "POL-2024-558812" in PdfReader(native).pages[0].extract_text()
    assert not PdfReader(scan).pages[0].extract_text().strip()
    presentation = Presentation(slides)
    native_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes
                            if shape.has_text_frame)
    assert "SIE-2025-001207" not in native_text
    assert any(shape.shape_type == 13 for slide in presentation.slides for shape in slide.shapes)


def test_image_slides_use_rendered_ocr_without_duplicate_native_text(tmp_path, monkeypatch):
    source = tmp_path / "slides.pptx"
    pptx_with_screenshot(source)
    monkeypatch.setattr(slide_render, "render", lambda _: pdf(2))
    calls = []

    def ocr(data, model, **kwargs):
        calls.append(model)
        assert len(PdfReader(io.BytesIO(data)).pages) == 1
        return {"segments": [{"text": "OCR includes SIE-2025-001207"}], "language": "es"}

    monkeypatch.setattr(docintel_client, "extract", ocr)
    result = pptx_extract.extract(source.read_bytes())
    assert calls == ["prebuilt-read", "prebuilt-read"]
    assert [s["location"] for s in result["segments"]] == ["slide:1", "slide:2"]
    assert all(s["text"] == "OCR includes SIE-2025-001207" for s in result["segments"])
