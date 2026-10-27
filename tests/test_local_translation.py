import io
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from PIL import Image
from pptx import Presentation

from scripts.local_demo import run_file
from services import local_translation
from services.config import ExternalServiceError
from tests.test_pipeline import pdf


def pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Polizza"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_sync_pptx_posts_document_and_glossary_without_storage(monkeypatch):
    seen = {}

    def request(path, job_id, **kwargs):
        seen.update(path=path, job_id=job_id, **kwargs)
        return SimpleNamespace(content=pptx())

    monkeypatch.setattr(local_translation, "request", request)
    output, usage = local_translation.translate_pptx(pptx(), b"polizza\tpolicy\n", "it", "en", "job")
    assert output.startswith(b"PK") and usage["glossaryApplied"]
    assert seen["path"] == "/translator/document:translate" and seen["job_id"] == "job"
    assert seen["params"]["api-version"] == "2024-05-01"
    assert set(seen["files"]) == {"document", "glossary"}
    assert usage["charactersCharged"] is None


def test_text_translation_splits_without_loss_and_preserves_locations(monkeypatch):
    calls = []

    def request(path, job_id, **kwargs):
        text = kwargs["json"][0]["Text"]
        calls.append(text)
        return SimpleNamespace(
            headers={"X-Metered-Usage": str(len(text))},
            json=lambda: [{"translations": [{"to": "en", "text": text}]}],
        )

    monkeypatch.setattr(local_translation, "request", request)
    text = "a " * 12000
    result, usage = local_translation.translate_segments(
        [{"text": text, "location": "paragraph:1"}], "it", "en", "job",
    )
    assert len(calls) == 3 and "".join(calls) == text
    assert result["segments"][0] == {"text": text, "location": "paragraph:1"}
    assert usage["charactersCharged"] == len(text) and not usage["glossaryApplied"]


def test_text_missing_meter_is_unknown_not_free(monkeypatch):
    monkeypatch.setattr(local_translation, "request", lambda *_a, **_kw: SimpleNamespace(
        headers={}, json=lambda: [{"translations": [{"to": "en", "text": "Policy"}]}],
    ))
    _, usage = local_translation.translate_segments([{"text": "Polizza"}], "it", "en", "job")
    assert usage["charactersCharged"] is None


@pytest.mark.parametrize("body", [[], [None], [{"translations": None}], [{"translations": [None]}],
                                   [{"translations": []}],
                                   [{"translations": [{"to": "it", "text": "wrong"}]}]])
def test_invalid_translation_is_not_success(monkeypatch, body):
    monkeypatch.setattr(local_translation, "request", lambda *_a, **_kw: SimpleNamespace(
        json=lambda: body,
    ))
    with pytest.raises(ExternalServiceError):
        local_translation.translate_segments([{"text": "Polizza"}], "it", "en", "job")


def test_keyless_request_does_not_follow_redirect_or_expose_error_body(monkeypatch):
    monkeypatch.setenv("TRANSLATOR_ENDPOINT", "https://unit.cognitiveservices.azure.com")
    identity = MagicMock()
    identity.get_token.return_value.token = "synthetic-token"
    monkeypatch.setattr(local_translation, "credential", lambda: identity)
    post = MagicMock(return_value=SimpleNamespace(status_code=401, text="PRIVATE CONTENT"))
    monkeypatch.setattr(local_translation.requests, "post", post)
    with pytest.raises(ExternalServiceError, match="HTTP 401; jobId=job") as error:
        local_translation.request("/translator/document:translate", "job")
    assert "PRIVATE CONTENT" not in str(error.value)
    assert not post.call_args.kwargs["allow_redirects"]
    assert post.call_args.kwargs["headers"]["Authorization"].startswith("Bearer ")


def test_sync_size_limit_is_enforced_before_network(monkeypatch):
    request = MagicMock()
    monkeypatch.setattr(local_translation, "request", request)
    with pytest.raises(ValueError, match="10 MB"):
        local_translation.translate_pptx(b"x" * (10 * 1024 * 1024 + 1), b"", "it", "en", "job")
    with pytest.raises(ValueError, match="1 MB"):
        local_translation.translate_pptx(b"", b"x" * (1024 * 1024 + 1), "it", "en", "job")
    request.assert_not_called()


def test_native_pptx_rejects_images_before_translation(monkeypatch):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, 0, 0)
    output = io.BytesIO()
    presentation.save(output)
    request = MagicMock()
    monkeypatch.setattr(local_translation, "request", request)
    with pytest.raises(ValueError, match="native text only"):
        local_translation.translate_pptx(output.getvalue(), b"", "it", "en", "job")
    request.assert_not_called()


def test_native_pptx_rejects_images_in_inherited_layouts():
    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(layout)
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(image, format="PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, 0, 0)
    image_relation = slide.part.rels[picture._element.blipFill.blip.rEmbed]
    layout_relation = layout.part.relate_to(image_relation.target_part, image_relation.reltype)
    picture._element.blipFill.blip.rEmbed = layout_relation
    slide.shapes._spTree.remove(picture._element)
    layout.shapes._spTree.insert_element_before(picture._element, "p:extLst")
    output = io.BytesIO()
    presentation.save(output)
    with pytest.raises(ValueError, match="native text only"):
        local_translation.native_pptx_only(output.getvalue())


def test_native_pptx_rejects_picture_placeholders():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[8])
    output = io.BytesIO()
    presentation.save(output)
    with pytest.raises(ValueError, match="native text only"):
        local_translation.native_pptx_only(output.getvalue())


def test_cli_failure_reports_only_exception_type_and_job_id(monkeypatch, capsys):
    from scripts import local_demo

    monkeypatch.setattr("sys.argv", ["local_demo.py", "claim.pdf"])
    monkeypatch.setattr(local_demo, "run_file", lambda *_args, **_kwargs: (_ for _ in ()).throw(
        ValueError("PRIVATE DOCUMENT EXCERPT")))
    assert local_demo.main() == 1
    output = capsys.readouterr()
    assert "ValueError" in output.err and "jobId=" in output.err
    assert "PRIVATE DOCUMENT EXCERPT" not in output.err
    assert "Traceback" not in output.err


def test_pdf_workflow_writes_local_report_and_flags_output(monkeypatch, tmp_path):
    from scripts import local_demo
    from services import storage
    for name in ("put", "get_json", "put_json", "source", "target"):
        monkeypatch.setattr(storage, name, MagicMock(side_effect=AssertionError("No cloud storage allowed")))
    source = tmp_path / "source.pdf"
    source.write_bytes(pdf())
    segments = [{"text": "Polizza", "role": "paragraph", "location": "paragraph:1"}]
    monkeypatch.setattr(local_demo.docintel_client, "extract", lambda *_a, **_kw: {
        "segments": segments, "language": "it", "pageCount": 1,
    })
    monkeypatch.setattr(local_demo, "prepare_terms", lambda extracted, *_: {**extracted, "terms": []})
    monkeypatch.setattr(local_demo.local_translation, "translate_segments", lambda *_: (
        {"segments": segments}, {"mode": "ocr-then-text", "charactersCharged": None},
    ))
    monkeypatch.setattr(local_demo, "review_documents", lambda job, *_: {
        "jobId": job["jobId"], "decision": "flagged", "score": 80, "issues": [], "costSignals": {},
    })
    stages = []
    report = run_file(source, tmp_path / "results", job_id="a" * 32, progress=stages.append)
    folder = tmp_path / "results" / report["jobId"]
    assert (folder / "report.json").is_file()
    assert (folder / "translated.txt").read_text().startswith("Machine translation - requires review")
    assert report["status"] == "needs_review" and not report["layoutPreserved"]
    assert report["costSignals"]["ocrPages"] == 1
    import json
    preview = json.loads((folder / "showcase.json").read_text(encoding="utf-8"))
    assert preview["sourceText"] == "Polizza" and preview["terms"] == []
    assert stages == ["validate", "extract", "prepare", "translate", "check", "save"]
    assert report["jobId"] == "a" * 32


@pytest.mark.parametrize("extension", ["pdf", "pptx"])
def test_translator_only_has_no_glossary_or_foundry_calls(monkeypatch, tmp_path, extension):
    from scripts import local_demo
    data = pdf() if extension == "pdf" else pptx()
    source = tmp_path / f"source.{extension}"
    source.write_bytes(data)
    extracted = {"segments": [{"text": "Polizza", "location": "paragraph:1"}], "language": "it", "pageCount": 1}
    monkeypatch.setattr(local_demo.docintel_client, "extract", lambda *_a, **_kw: extracted)
    forbidden = MagicMock(side_effect=AssertionError("Baseline must not call agents"))
    monkeypatch.setattr(local_demo, "prepare_terms", forbidden)
    monkeypatch.setattr(local_demo, "review_documents", forbidden)
    monkeypatch.setattr(local_demo.local_translation, "translate_segments", lambda *_: (
        {"segments": [{"text": "Policy"}]}, {"glossaryApplied": False, "charactersCharged": 7},
    ))
    def translate_pptx(data, glossary, *_):
        assert glossary == b""
        return pptx(), {"glossaryApplied": False, "charactersCharged": None}
    monkeypatch.setattr(local_demo.local_translation, "translate_pptx", translate_pptx)
    stages = []
    report = run_file(source, tmp_path / "out", translator_only=True, progress=stages.append)
    assert report["status"] == "not_reviewed" and report["score"] is None
    assert report["workflow"] == "translator_only" and report["costSignals"]["tokens"] == 0
    import hashlib
    assert report["sourceFile"] == source.name
    assert report["sourceSha256"] == hashlib.sha256(data).hexdigest()
    assert "prepare" not in stages and "check" not in stages
    assert report["passB"] == [] and report["versions"] == []
    forbidden.assert_not_called()
